from pathlib import Path

import pytest

from parser.document_loader import load_document
from parser.extractor import ResumeExtractor
from parser.errors import UnsupportedFileTypeError


def test_load_txt(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Jane Doe\n", encoding="utf-8")
    assert load_document(path) == "Jane Doe\n"


def test_unsupported_file_type(tmp_path: Path) -> None:
    path = tmp_path / "resume.md"
    path.write_text("Jane", encoding="utf-8")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(path)


def test_load_docx_with_table_only_resume_content(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "table_resume.docx"
    document = Document()
    table = document.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "Jane Doe"
    table.cell(0, 1).text = "jane@example.com"
    table.cell(1, 0).text = "Skills"
    table.cell(1, 1).text = "Python, SQL"
    table.cell(2, 0).text = "Experience"
    table.cell(2, 1).text = "Presentation Designer at Acme | 2020 - 2024"
    table.cell(3, 0).text = "Education"
    table.cell(3, 1).text = "State University 2016 - 2020"
    document.save(path)

    text = load_document(path)

    assert "Jane Doe" in text
    assert "Presentation Designer at Acme | 2020 - 2024" in text
    assert "  " not in text

    resume = ResumeExtractor().extract(text)
    assert resume.name == "Jane Doe"
    assert resume.email == "jane@example.com"
    assert resume.work_experience[0].title == "Presentation Designer"
    assert resume.work_experience[0].company == "Acme"
    assert resume.work_experience[0].dates == "2020 - 2024"
